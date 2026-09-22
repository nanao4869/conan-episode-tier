#!/usr/bin/env python3
# Merges the user's personal Excel rating sheet ("名探偵コナン アニメ評価") into their existing exported Tier
# board, so the インポート button can load the result directly.
#
#   python tools/ratings-to-import.py <current board .json> <ratings .xlsx> <output .json path>
#
# The current board is kept as-is (its rows, names, colors and whatever is already placed in them);
# this script only ADDS items that aren't already placed anywhere in it. Star ratings 4/3/2/1 go into the
# rows the user named for them (see STAR_TO_ROW_NAME below - edit it if their row names differ). Star 5
# is NOT auto-placed (the user ranks 5-star episodes among themselves by hand into their own S/A/B/C/D
# rows) - this script only reports how many unplaced 5-star episodes there are. Every matched episode's
# 感想 is saved as that item's memo regardless of whether it was placed (see [[project_conan_episode_tier]]
# エピソードのメモ機能), capped at 300 chars like the app.
#
# Expected sheet 1 layout (header row 2): シーズン, 話数 (cumulative TV episode number - the same numbering
# YTV/Wikipedia use, which is what data/episodes.js is built from), タイトル, 評価 (star string, e.g.
# "★★★★☆"), 感想, 放送日, 種別 (原作/アニメ). Rows with no 評価 are skipped (not yet rated).
#
# A TV episode belongs to exactly one Tier item: either a manga case (its `an[]` entries list which TV
# episode(s) adapted it) or an anime-original entry (its `file` text has the same "第N話"/"第N・M話（全2話）"/
# "第N〜M話（全K話）" range). Several TV episodes can share one item (multi-part episodes, or a manga case
# animated twice) - their stars are averaged (rounded to the nearest star) and their 感想 are joined.
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
NOTE_MAX = 300

# star count -> the name of the row it goes into (must match a tier's "name" in the current board exactly)
STAR_TO_ROW_NAME = {4: "E", 3: "F", 2: "G", 1: "H"}


def load_episodes():
    text = (ROOT / "data" / "episodes.js").read_text(encoding="utf-8")
    m = re.search(r"window\.EPISODES = (\[.*\]);", text, re.S)
    return json.loads(m.group(1))


def parse_range(text):
    """"第13話" -> [13]; "第27・28話（全2話）" -> [27,28]; "第68〜70話（全3話）" -> [68,69,70]."""
    m = re.search(r"第([0-9・〜]+)話", text)
    if not m:
        return []
    body = m.group(1)
    if "〜" in body:
        a, b = body.split("〜")
        return list(range(int(a), int(b) + 1))
    return [int(x) for x in body.split("・")]


def build_ep_index(episodes):
    """TV episode number -> (item no, kind)."""
    idx = {}
    for e in episodes:
        if e.get("k") == "a":
            for ep in parse_range(e["file"]):
                idx[ep] = (e["no"], "anime")
        else:
            for a in e.get("an") or []:
                for ep in parse_range(a.get("e", "")):
                    idx[ep] = (e["no"], "manga")
    return idx


def main():
    if len(sys.argv) != 4:
        print("usage: python ratings-to-import.py <current board .json> <ratings .xlsx> <output .json path>")
        sys.exit(1)
    board_path, xlsx_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]

    board = json.loads(Path(board_path).read_text(encoding="utf-8"))
    tiers = board["tiers"]
    row_by_name = {t["name"]: t for t in tiers}
    for star, name in STAR_TO_ROW_NAME.items():
        if name not in row_by_name:
            print(f"! row \"{name}\" (for {star}-star episodes) does not exist in {board_path} - check STAR_TO_ROW_NAME")
            sys.exit(1)
    already_placed = {n for t in tiers for n in t["items"]}

    episodes = load_episodes()
    ep_index = build_ep_index(episodes)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    by_item = {}
    unmatched = []
    for r in range(3, ws.max_row + 1):
        ep, title, stars_s, comment, kind = (ws.cell(r, c).value for c in (2, 3, 4, 5, 7))
        if not stars_s:
            continue
        ep = int(ep)
        stars = stars_s.count("★")
        hit = ep_index.get(ep)
        if not hit:
            unmatched.append((ep, title))
            continue
        item_no, _ = hit
        slot = by_item.setdefault(item_no, {"stars": [], "comments": []})
        slot["stars"].append(stars)
        if comment:
            slot["comments"].append(comment.strip())

    notes = dict(board.get("notes") or {})
    added = {name: 0 for name in STAR_TO_ROW_NAME.values()}
    skipped_placed = {name: 0 for name in STAR_TO_ROW_NAME.values()}
    five_star_new = 0
    five_star_already_placed = 0

    for item_no, slot in sorted(by_item.items()):
        avg = round(sum(slot["stars"]) / len(slot["stars"]))
        avg = max(1, min(5, avg))

        if slot["comments"]:
            text = "／".join(slot["comments"])
            if len(text) > NOTE_MAX:
                text = text[: NOTE_MAX - 1] + "…"
            notes[str(item_no)] = text

        if avg == 5:
            if item_no in already_placed:
                five_star_already_placed += 1
            else:
                five_star_new += 1
            continue

        row_name = STAR_TO_ROW_NAME[avg]
        if item_no in already_placed:
            skipped_placed[row_name] += 1
            continue
        row_by_name[row_name]["items"].append(item_no)
        already_placed.add(item_no)
        added[row_name] += 1

    board["notes"] = notes
    Path(out_path).write_text(json.dumps(board, ensure_ascii=False, indent=2), encoding="utf-8")

    print("added to each row (items that were not already placed anywhere):")
    for star, name in STAR_TO_ROW_NAME.items():
        print(f"  {name} (星{star}): +{added[name]}  (already placed elsewhere, left alone: {skipped_placed[name]})")
    print(f"星5: {five_star_new} 件が未配置のまま残っています（{five_star_already_placed} 件は既にS〜Dのどこかに配置済み）")
    print(f"memos saved: {len(notes)} (includes any that were already in the board)")
    print(f"unmatched rated rows (TV episode number not found in data/episodes.js): {len(unmatched)}")
    for ep, title in unmatched[:30]:
        print(f"    ep{ep} {title}")


if __name__ == "__main__":
    main()
